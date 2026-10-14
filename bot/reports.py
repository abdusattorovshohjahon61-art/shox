"""Excel hisobot va har kungi avtomatik hisobot."""
import asyncio
import io
import logging
from datetime import datetime, timedelta

from aiogram import Bot
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from bot.config import Config
from bot.db import Database
from bot.handlers.common import notify_admins
from bot.i18n import status_name
from bot.utils import escape, money

log = logging.getLogger(__name__)


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, name in enumerate(header, 1):
        width = max([len(str(name))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)


def build_excel(db: Database) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    orders = db.get_orders(limit=None)
    _sheet(
        wb, "Buyurtmalar",
        ["№", "Sana", "Holat", "Mijoz", "Telefon", "Manzil", "Qabul qilish", "To'lov", "To'langan",
         "Mahsulotlar", "Summa", "Chegirma", "Yetkazish", "Jami", "Promokod", "Baho", "Izoh"],
        [
            [
                o["id"], o["created_at"], status_name(o["status"]), o["full_name"], o["phone"], o["address"],
                "Olib ketish" if o["delivery_type"] == "pickup" else "Yetkazib berish",
                "Karta" if o["payment_method"] == "card" else "Naqd",
                "Ha" if o["paid"] else "Yo'q",
                "; ".join(f"{i['name']} x{i['qty']}" for i in db.get_order_items(o["id"])),
                o["subtotal"], o["discount"], o["delivery_fee"], o["total"], o["promo_code"] or "",
                o["rating"] or "", o["comment"],
            ]
            for o in orders
        ],
    )

    products = [p for c in db.get_categories() for p in db.get_products(c["id"], include_hidden=True)]
    categories = {c["id"]: c["name"] for c in db.get_categories()}
    _sheet(
        wb, "Mahsulotlar",
        ["ID", "Kategoriya", "Nomi", "Narxi", "Qoldiq", "Holati"],
        [
            [p["id"], categories.get(p["category_id"], ""), p["name"], p["price"],
             "∞" if p["stock"] is None else p["stock"], "Sotuvda" if p["is_active"] else "Yashirin"]
            for p in products
        ],
    )

    _sheet(
        wb, "Mijozlar",
        ["Telegram ID", "Ism", "Username", "Telefon", "Ro'yxatdan o'tgan", "Buyurtmalar", "Jami xarid"],
        [
            [u["id"], u["full_name"], f"@{u['username']}" if u["username"] else "", u["phone"] or "",
             u["created_at"], u["orders_count"], u["spent"]]
            for u in db.get_customers()
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stats_text(db: Database, title: str, day: str | None = None) -> str:
    s = db.stats(day)
    lines = [
        f"📊 <b>{title}</b>",
        "",
        f"🧾 Buyurtmalar: {s['orders']} (yangi: {s['new_orders']}, bekor: {s['cancelled']})",
        f"💵 Savdo (bekor qilinganlarsiz): {money(s['sales'])}",
        f"✅ Yetkazilganlar summasi: {money(s['revenue'])}",
    ]
    if s["avg_rating"]:
        lines.append(f"⭐️ O'rtacha baho: {s['avg_rating']:.1f}")
    if not day:
        lines.append(f"👥 Foydalanuvchilar: {s['users']}")
        lines.append(f"📅 Bugungi buyurtmalar: {s['today_orders']}")
    if s["top_products"]:
        lines += ["", "🔥 <b>Eng ko'p sotilganlar:</b>"]
        lines += [f"{i}. {escape(name)} — {qty} dona" for i, (name, qty) in enumerate(s["top_products"], 1)]
    low = db.low_stock_products()
    if low:
        lines += ["", "📉 <b>Omborda kam qolgan:</b>"]
        lines += [f"• {escape(p['name'])} — {p['stock']} dona" for p in low[:10]]
    return "\n".join(lines)


def seconds_until(target, now: datetime) -> float:
    run_at = now.replace(hour=target.hour, minute=target.minute, second=0, microsecond=0)
    if run_at <= now:
        run_at += timedelta(days=1)
    return (run_at - now).total_seconds()


async def daily_report_loop(bot: Bot, db: Database, config: Config) -> None:
    """Har kuni config.report_time da adminlarga kunlik hisobot yuboradi."""
    if not config.report_time or not config.admin_ids:
        return
    while True:
        await asyncio.sleep(seconds_until(config.report_time, config.now()))
        try:
            today = config.now().strftime("%Y-%m-%d")
            await notify_admins(bot, config, stats_text(db, f"Kunlik hisobot — {today}", day=today))
        except Exception:
            log.exception("Kunlik hisobot yuborilmadi")
