import asyncio
from datetime import datetime, time

import pytest
from aiogram.fsm.storage.base import StorageKey

from bot.config import Config
from bot.db import Database, OutOfStock
from bot.reports import build_excel, seconds_until, stats_text
from bot.storage import SQLiteStorage
from bot.utils import money, normalize_phone, parse_int


@pytest.fixture
def db():
    d = Database(":memory:")
    yield d
    d.close()


@pytest.fixture
def shop(db):
    cat = db.add_category("Test")
    a = db.add_product(cat, "Olma", 10_000)
    b = db.add_product(cat, "Nok", 5_000, stock=2)
    return cat, a, b


def test_seed_is_idempotent(db):
    db.seed_demo()
    db.seed_demo()
    assert len(db.get_categories()) == 3
    assert db.count_products() == 6


def test_cart_and_order_flow(db, shop):
    _, a, b = shop
    db.add_to_cart(1, a, 2)
    db.add_to_cart(1, a, 1)
    db.add_to_cart(1, b)
    assert {r["product_id"]: r["qty"] for r in db.get_cart(1)} == {a: 3, b: 1}
    assert db.cart_total(1) == 35_000

    db.set_cart_qty(1, b, 0)
    assert db.cart_total(1) == 30_000

    order_id = db.create_order(1, "Ali", "+998901234567", "Toshkent")
    assert db.get_cart(1) == []
    order = db.get_order(order_id)
    # 30 000 + 15 000 yetkazib berish (standart sozlama)
    assert (order["subtotal"], order["delivery_fee"], order["total"]) == (30_000, 15_000, 45_000)
    assert [(i["name"], i["qty"]) for i in db.get_order_items(order_id)] == [("Olma", 3)]

    # Narx o'zgarsa ham buyurtmadagi narx saqlanib qoladi
    db.update_product_price(a, 99_000)
    assert db.get_order_items(order_id)[0]["price"] == 10_000

    db.set_order_status(order_id, "done")
    assert db.stats()["revenue"] == 45_000
    with pytest.raises(ValueError):
        db.set_order_status(order_id, "xato")


def test_quote_pickup_free_delivery_and_promo(db, shop):
    _, a, _ = shop
    db.add_to_cart(1, a, 40)  # 400 000
    db.add_promo("vip", 10, max_uses=1)

    assert db.quote(1, "pickup").delivery_fee == 0
    q = db.quote(1, "delivery", "VIP")
    assert (q.subtotal, q.discount, q.delivery_fee, q.total) == (400_000, 40_000, 0, 360_000)

    db.set_setting("free_delivery_from", 0)  # bepul yetkazish o'chirildi
    assert db.quote(1, "delivery").delivery_fee == 15_000

    db.create_order(1, "Ali", "+998901234567", "Toshkent", promo_code="vip")
    assert db.get_promo("VIP") is None  # limit tugadi


def test_stock_is_enforced_and_restored(db, shop):
    _, _, b = shop
    db.add_to_cart(1, b, 3)
    with pytest.raises(OutOfStock) as e:
        db.create_order(1, "Ali", "+998901234567", "Toshkent")
    assert e.value.items == [("Nok", 2)]

    db.set_cart_qty(1, b, 2)
    oid = db.create_order(1, "Ali", "+998901234567", "Toshkent")
    assert db.get_product(b)["stock"] == 0
    assert [p["id"] for p in db.low_stock_products()] == [b]

    db.set_order_status(oid, "cancelled")
    assert db.get_product(b)["stock"] == 2
    db.set_order_status(oid, "accepted")  # qayta tiklash yana ayiradi
    assert db.get_product(b)["stock"] == 0


def test_hidden_product_leaves_cart_and_catalog(db):
    cat = db.add_category("Test")
    p = db.add_product(cat, "Olma", 1000)
    db.add_to_cart(1, p)
    db.update_product(p, is_active=0)
    assert db.get_cart(1) == []
    assert db.get_products(cat) == []
    assert len(db.get_products(cat, include_hidden=True)) == 1
    assert db.get_categories(only_with_products=True) == []
    with pytest.raises(ValueError):
        db.update_product(p, id=5)


def test_repeat_order(db, shop):
    _, a, b = shop
    db.add_to_cart(1, a, 2)
    db.add_to_cart(1, b, 1)
    oid = db.create_order(1, "Ali", "+998901234567", "Toshkent")
    db.delete_product(b)
    assert db.refill_cart_from_order(1, oid) == 1
    assert [(r["name"], r["qty"]) for r in db.get_cart(1)] == [("Olma", 2)]


def test_search_is_case_insensitive_for_cyrillic(db):
    cat = db.add_category("Test")
    db.add_product(cat, "Шоколад", 1000)
    db.add_product(cat, "Choy", 1000, "Ko'k choy")
    assert [r["name"] for r in db.search_products("шоколад")] == ["Шоколад"]
    assert [r["name"] for r in db.search_products("KO'K")] == ["Choy"]


def test_migration_adds_missing_columns(tmp_path):
    import sqlite3
    path = tmp_path / "old.db"
    conn = sqlite3.connect(path)
    conn.executescript(
        "CREATE TABLE products (id INTEGER PRIMARY KEY, category_id INTEGER, name TEXT, "
        "description TEXT DEFAULT '', price INTEGER, photo_id TEXT, is_active INTEGER DEFAULT 1);"
    )
    conn.close()
    db = Database(str(path))
    cols = {r["name"] for r in db.conn.execute("PRAGMA table_info(products)")}
    assert "stock" in cols
    db.close()


def test_sqlite_fsm_storage_survives_restart(tmp_path):
    path = str(tmp_path / "s.db")
    key = StorageKey(bot_id=1, chat_id=2, user_id=2)

    async def scenario():
        db = Database(path)
        s = SQLiteStorage(db)
        await s.set_state(key, "Checkout:phone")
        await s.set_data(key, {"name": "Ali"})
        db.close()

        db2 = Database(path)  # "qayta ishga tushirish"
        s2 = SQLiteStorage(db2)
        assert await s2.get_state(key) == "Checkout:phone"
        assert await s2.get_data(key) == {"name": "Ali"}
        await s2.set_state(key, None)
        assert await s2.get_state(key) is None
        db2.close()

    asyncio.run(scenario())


def test_excel_and_stats(db, shop):
    from io import BytesIO
    from openpyxl import load_workbook

    _, a, _ = shop
    db.upsert_user(1, "Ali", "ali")
    db.add_to_cart(1, a, 2)
    db.create_order(1, "Ali", "+998901234567", "Toshkent")
    wb = load_workbook(BytesIO(build_excel(db)))
    assert wb.sheetnames == ["Buyurtmalar", "Mahsulotlar", "Mijozlar"]
    assert wb["Buyurtmalar"].max_row == 2
    assert wb["Mijozlar"]["F2"].value == 1
    assert "Olma — 2 dona" in stats_text(db, "Test")


def test_working_hours():
    day = Config(bot_token="x", open_time=time(9), close_time=time(21))
    assert day.is_open(datetime(2026, 1, 1, 10, 0))
    assert not day.is_open(datetime(2026, 1, 1, 22, 0))
    night = Config(bot_token="x", open_time=time(20), close_time=time(4))
    assert night.is_open(datetime(2026, 1, 1, 2, 0))
    assert not night.is_open(datetime(2026, 1, 1, 12, 0))
    assert Config(bot_token="x").is_open()


def test_seconds_until():
    now = datetime(2026, 1, 1, 21, 0)
    assert seconds_until(time(21, 30), now) == 30 * 60
    assert seconds_until(time(20, 0), now) == 23 * 3600


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("+998 90 123-45-67", "+998901234567"),
        ("998901234567", "+998901234567"),
        ("901234567", "+998901234567"),
        ("12345", None),
        ("+7 999 123 45 67", None),
    ],
)
def test_normalize_phone(raw, expected):
    assert normalize_phone(raw) == expected


def test_money_and_parse_int():
    assert money(1250000) == "1 250 000 so'm"
    assert money(1000, "ru") == "1 000 сум"
    assert parse_int("125 000 so'm") == 125000
    assert parse_int("abc") is None
